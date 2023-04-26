# -*- coding: utf-8 -*-
"""

"""

import pandas as pd
import numpy as np
import docx
from docx import Document

doc_filepath='机场四字码.docx'

city_airport_code={'城市':[],'机场四字码':[]}
document=Document(doc_filepath)
tables=document.tables
table=tables[0]
city_airport_code['城市']=[table.cell(i,0).text for i in range(1,len(table.rows))]
city_airport_code['机场四字码']=[table.cell(i,3).text for i in range(1,len(table.rows))]

def is_chinese(string):
    for ch in string:
        if u'\u4e00' <= ch <= u'\u9fff':
            return True
def get_chinese(string):
    ch_str=''
    for ch in string:
        if u'\u4e00' <= ch <= u'\u9fff':
            ch_str+=ch
    return ch_str

city_code=list(zip(city_airport_code['城市'],city_airport_code['机场四字码']))
city_code=[tuple_data for tuple_data in city_code if tuple_data[1]!=''  and is_chinese(tuple_data[0])]
city_code=[(get_chinese(city_code[i][0]),city_code[i][1]) for i in range(len(city_code)) if get_chinese(city_code[i][0]) not in ['香港','台北','澳门','高雄','台东','金门','马祖','台中','马公']]
code_map={code[1]:code[0] for code in city_code}

city=list(set([city_code[i][0]  for i in range(len(city_code))]))
airline_info=pd.read_excel('0701flt.xls')
timeable=airline_info.loc[:,'P_DEPAP':].dropna()


departure=list(timeable['P_DEPAP'])
arrive=list(timeable['P_ARRAP'])

timeable['出发城市']=[code_map[code]  if code in code_map.keys() else None for code in departure ]
timeable['到达城市']=[code_map[code]  if code in code_map.keys() else None for code in arrive]
timeable=timeable.iloc[:,2:].dropna()






