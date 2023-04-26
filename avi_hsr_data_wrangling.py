# -*- coding: utf-8 -*-
"""

"""

import pandas as pd
import numpy as np
import docx
from docx import Document
from datetime import datetime

doc_filepath='机场四字码.docx'
city_airport_code={'城市':[],'机场四字码':[]}
document=Document(doc_filepath)
tables=document.tables
table=tables[0]
city_airport_code['城市']=[table.cell(i,0).text for i in range(1,len(table.rows))]
city_airport_code['机场四字码']=[table.cell(i,3).text for i in range(1,len(table.rows))]

def is_chinese(string):
    for ch in string:
        if u'\u4e00' <= ch <= u'\u9fff':
            return True
def get_chinese(string):
    ch_str=''
    for ch in string:
        if u'\u4e00' <= ch <= u'\u9fff':
            ch_str+=ch
    return ch_str

city_code=list(zip(city_airport_code['城市'],city_airport_code['机场四字码']))
city_code=[tuple_data for tuple_data in city_code if tuple_data[1]!=''  and is_chinese(tuple_data[0])]
city_code=[(get_chinese(city_code[i][0]),city_code[i][1]) for i in range(len(city_code)) if get_chinese(city_code[i][0]) not in ['香港','台北','澳门','高雄','台东','金门','马祖','台中','马公']]
code_map={code[1]:code[0] for code in city_code}

city=list(set([city_code[i][0]  for i in range(len(city_code))]))
airline_info=pd.read_excel('0701flt.xls')
del airline_info['P_AIRCRAFTTYPE']
del airline_info['P_REGISTENUM']
timeable=airline_info.loc[:,'FLIGHTID':].dropna()
timeable.index=pd.Index(range(timeable.shape[0]))




departure=list(timeable['P_DEPAP'])
arrive=list(timeable['P_ARRAP'])

timeable['出发城市']=[code_map[code]  if code in code_map.keys() else None for code in departure ]
timeable['到达城市']=[code_map[code]  if code in code_map.keys() else None for code in arrive]
timeable=timeable.loc[:,['FLIGHTID','出发城市','到达城市','P_DEPTIME','P_ARRTIME']].dropna()
timeable.rename(columns={'P_DEPTIME':"departure_time",'P_ARRTIME':'arrival_time'},inplace='True')  

time_str=lambda x:datetime.strptime(str(int(x))[-4:-2]+":"+str(int(x))[-2:],"%H:%M")
timeable['departure_time']=timeable['departure_time'].map(time_str)
timeable['arrival_time']=timeable['arrival_time'].map(time_str)
timeable['airlne_timecost']=timeable['arrival_time']-timeable['departure_time']

timeable.to_excel('Airline_timeable.xlsx',index=False)




railway_info=pd.read_excel('crawl_train_info.xlsx')
railway_info_df=railway_info.fillna(method='ffill')
f=lambda x:None if x=='[]' else x
for column in railway_info_df.columns:
    railway_info_df[column]=railway_info_df[column].map(f)  
railway_info_df_clean=railway_info_df.dropna()
railway_info_df_clean.rename(columns={'Unnamed: 0':"出发城市",'Unnamed: 1':'到达城市','train_num':'TrainID'},inplace='True')   

for column in railway_info_df_clean.columns[2:]:
    railway_info_df_clean[column]=railway_info_df_clean[column].map(lambda x: x.split('[')[1].split(']')[0])

transport_type={'train':0,'airplane':1}

expand_railway_dict={'出发城市':[],'到达城市':[],'departure_time':[],'arrival_time':[],'departure_station':[],'arrival_station':[],'travel_time':[],'TrainID':[]}
for index in railway_info_df_clean.index:
      if len(railway_info_df_clean.loc[index,'departure_time'].split(','))==1:
          for column in expand_railway_dict.keys():
            expand_railway_dict[column].append(railway_info_df_clean.loc[index,column])
      else:
            expand_railway_dict['出发城市']+=[railway_info_df_clean.loc[index,'出发城市']]* len(railway_info_df_clean.loc[index,'departure_time'].split(','))
            expand_railway_dict['到达城市']+=[railway_info_df_clean.loc[index,'到达城市']]* len(railway_info_df_clean.loc[index,'departure_time'].split(','))
            for column in expand_railway_dict.keys():
                if column not in ['出发城市','到达城市']:
                    expand_railway_dict[column]+=railway_info_df_clean.loc[index,column].split(',')
                    
railway_df=pd.DataFrame(expand_railway_dict)
timestr=lambda x:datetime.strptime(x.split('\'')[1],"%H:%M")
timestr2=lambda x:datetime.strptime(x.split('\'')[1].strip(),"%H:%M")
f=lambda x:None if x[1]!='G' else x
railway_df['TrainID']=railway_df['TrainID'].map(f)
railway_df_clean=railway_df.dropna(axis=0,how='any')
for column in ['departure_station','arrival_station','TrainID']:
     railway_df_clean[column]=railway_df_clean[column].map(lambda x:x.split('\'')[1].strip())
railway_df_clean['departure_time']=railway_df_clean['departure_time'].map(timestr)
railway_df_clean['arrival_time']=railway_df_clean['arrival_time'].map(timestr2)
railway_df_clean['travel_time']=railway_df_clean['arrival_time']-railway_df_clean['departure_time']
railway_df_clean.index=pd.Index(range(railway_df_clean.shape[0]))
railway_df_clean['transport_type']=pd.Series([transport_type['train']]*railway_df_clean.shape[0])
railway_df_clean.to_excel('High_speed_railway_info_cleaned.xlsx',index=False)









